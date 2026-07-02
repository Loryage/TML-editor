#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TML Markdown Editor 
主程序代码，现已支持多端同步
依赖见requirements.txt
"""

import sys
import os
import ctypes
import html as html_lib
import re
import importlib
import time
import traceback

# PyInstaller 隐藏导入标记：让静态分析能检测到懒加载的模块
# 运行时不会执行（if False），但 PyInstaller 会扫描到并打包这些依赖
if False:
    import markdown
    import markdown.extensions.extra
    import markdown.extensions.tables
    import markdown.extensions.fenced_code
    import docx
    import requests
    from PyQt6.QtWebEngineWidgets import QWebEngineView
    from PyQt6.QtWebEngineCore import QWebEngineSettings
    from PyQt6.QtWebChannel import QWebChannel

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QSplitter, QTextEdit,
    QFileDialog, QMessageBox, QMenuBar, QMenu,
    QStatusBar, QWidget, QVBoxLayout, QHBoxLayout,
    QTabWidget, QDialog, QPushButton, QSplashScreen,
    QLabel, QLineEdit, QFormLayout, QDialogButtonBox,
    QProgressBar, QCheckBox, QListWidget, QListWidgetItem,
    QAbstractItemView, QFrame
)
from PyQt6.QtCore import QThread, pyqtSignal
from PyQt6.QtGui import QAction, QFont, QSyntaxHighlighter, QTextCharFormat, QColor, QIcon, QTextCursor, QPixmap, QPainter, QLinearGradient
from PyQt6.QtCore import Qt, QFileInfo, QUrl, QObject, pyqtSignal, pyqtSlot, QTimer


_MARKDOWN_MODULE = None
_DOCX_DOCUMENT_CLASS = None
_DOCX_IMPORT_FAILED = False


# ==================== 启动日志模块 ====================
class StartupLogger:
    def __init__(self):
        self.started_at = time.perf_counter()
        self.log_path = self._resolve_log_path()
        self.enabled = False

    def _resolve_log_path(self):
        if getattr(sys, "frozen", False):
            base_dir = os.path.dirname(sys.executable)
        else:
            base_dir = os.path.dirname(os.path.abspath(__file__))
        return os.path.join(base_dir, "TMLEditor_startup.log")

    def _write(self, message):
        if not self.enabled:
            return
        elapsed_ms = (time.perf_counter() - self.started_at) * 1000.0
        line = f"[{elapsed_ms:9.1f} ms] {message}\n"
        try:
            with open(self.log_path, "a", encoding="utf-8") as log_file:
                log_file.write(line)
        except Exception:
            self.enabled = False

    def mark(self, message):
        self._write(message)

    def section(self, message):
        self._write(message)


STARTUP_LOGGER = StartupLogger()


# ==================== 同步工作线程 ====================
class SyncWorker(QThread):
    progress_signal = pyqtSignal(str, int)
    finished_signal = pyqtSignal(dict)

    def __init__(self, sync_client):
        super().__init__()
        self.sync_client = sync_client

    def run(self):
        def on_progress(message, percent):
            self.progress_signal.emit(message, percent)
        result = self.sync_client.sync(progress_callback=on_progress)
        self.finished_signal.emit(result)


def get_markdown_module():
    global _MARKDOWN_MODULE
    if _MARKDOWN_MODULE is None:
        _MARKDOWN_MODULE = importlib.import_module("markdown")
    return _MARKDOWN_MODULE


def get_docx_document_class():
    global _DOCX_DOCUMENT_CLASS, _DOCX_IMPORT_FAILED
    if _DOCX_DOCUMENT_CLASS is not None:
        return _DOCX_DOCUMENT_CLASS
    if _DOCX_IMPORT_FAILED:
        return None
    try:
        _DOCX_DOCUMENT_CLASS = importlib.import_module("docx").Document
        return _DOCX_DOCUMENT_CLASS
    except Exception:
        _DOCX_IMPORT_FAILED = True
        return None


def log_startup(message):
    STARTUP_LOGGER.mark(message)


def install_global_exception_hook():
    def hook(exc_type, exc_value, exc_traceback):
        try:
            log_startup("unhandled exception")
            log_startup("".join(traceback.format_exception(exc_type, exc_value, exc_traceback)).rstrip())
        finally:
            sys.__excepthook__(exc_type, exc_value, exc_traceback)

    sys.excepthook = hook


# ==================== 启动画面（Splash Screen） ====================
class SplashScreen(QSplashScreen):
    def __init__(self):
        pixmap = QPixmap(400, 280)
        pixmap.fill(QColor("#1e1e2e"))
        super().__init__(pixmap)
        self.setFixedSize(400, 280)
        self.setWindowFlag(Qt.WindowType.WindowStaysOnTopHint, True)

        self._progress = 0
        self._message = "正在启动..."
        self._dots = 0
        self._dot_timer = QTimer(self)
        self._dot_timer.timeout.connect(self._animate_dots)
        self._dot_timer.start(400)

        self._draw_content()

    def _animate_dots(self):
        self._dots = (self._dots + 1) % 4
        self._draw_content()

    def set_progress(self, value, message=None):
        self._progress = max(0, min(100, value))
        if message:
            self._message = message
        self._draw_content()

    def _draw_content(self):
        pixmap = QPixmap(400, 280)
        pixmap.fill(QColor("#1e1e2e"))
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing, True)

        title_font = QFont("Microsoft YaHei", 22, QFont.Weight.Bold)
        painter.setFont(title_font)
        painter.setPen(QColor("#89b4fa"))
        painter.drawText(pixmap.rect(), Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignTop,
                         "\n\nTML Markdown")

        subtitle_font = QFont("Microsoft YaHei", 10)
        painter.setFont(subtitle_font)
        painter.setPen(QColor("#a6adc8"))
        painter.drawText(pixmap.rect(), Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignTop,
                         "\n\n\n\n\n-")

        dots = "." * self._dots
        msg_font = QFont("Microsoft YaHei", 9)
        painter.setFont(msg_font)
        painter.setPen(QColor("#cdd6f4"))
        painter.drawText(pixmap.rect(), Qt.AlignmentFlag.AlignHCenter | Qt.AlignmentFlag.AlignBottom,
                         f"\n\n{self._message}{dots}\n\n")

        bar_x = 50
        bar_y = 190
        bar_w = 300
        bar_h = 6

        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(QColor("#313244"))
        painter.drawRoundedRect(bar_x, bar_y, bar_w, bar_h, 3, 3)

        fill_w = int(bar_w * (self._progress / 100.0))
        if fill_w > 0:
            gradient = QLinearGradient(bar_x, 0, bar_x + bar_w, 0)
            gradient.setColorAt(0, QColor("#89b4fa"))
            gradient.setColorAt(1, QColor("#cba6f7"))
            painter.setBrush(gradient)
            painter.drawRoundedRect(bar_x, bar_y, fill_w, bar_h, 3, 3)

        painter.end()
        self.setPixmap(pixmap)


# ==================== 配置与最近文件管理 ====================
def get_app_config_dir() -> str:
    if os.name == "nt":
        base = os.environ.get("APPDATA", os.path.expanduser("~"))
    else:
        base = os.path.expanduser("~/.config")
    app_dir = os.path.join(base, "TMLEditor")
    os.makedirs(app_dir, exist_ok=True)
    return app_dir


def get_settings_path() -> str:
    return os.path.join(get_app_config_dir(), "settings.json")


def load_settings() -> dict:
    path = get_settings_path()
    default = {
        "show_quick_open_on_start": True,
        "recent_files": [],
    }
    if os.path.exists(path):
        try:
            import json
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
            for k, v in default.items():
                if k not in data:
                    data[k] = v
            return data
        except Exception:
            pass
    return dict(default)


def save_settings(settings: dict):
    path = get_settings_path()
    try:
        import json
        with open(path, "w", encoding="utf-8") as f:
            json.dump(settings, f, indent=2, ensure_ascii=False)
    except Exception:
        pass


def get_recent_files() -> list:
    settings = load_settings()
    return settings.get("recent_files", [])


def add_recent_file(file_path: str):
    settings = load_settings()
    recent = settings.get("recent_files", [])
    recent = [r for r in recent if r.get("path") != file_path]
    recent.insert(0, {
        "path": file_path,
        "name": os.path.basename(file_path),
        "time": int(time.time())
    })
    recent = recent[:20]
    settings["recent_files"] = recent
    save_settings(settings)


def remove_recent_file(file_path: str):
    settings = load_settings()
    recent = [r for r in settings.get("recent_files", []) if r.get("path") != file_path]
    settings["recent_files"] = recent
    save_settings(settings)


# ==================== 快速打开对话框 ====================
class QuickOpenDialog(QDialog):
    def __init__(self, sync_folder: str = "", parent=None):
        super().__init__(parent)
        self.setWindowTitle("快速打开")
        self.resize(520, 420)
        self.selected_path = None

        layout = QVBoxLayout(self)
        layout.setSpacing(12)

        title = QLabel("选择要打开的文件")
        title_font = QFont()
        title_font.setPointSize(12)
        title_font.setBold(True)
        title.setFont(title_font)
        layout.addWidget(title)

        content_layout = QHBoxLayout()
        content_layout.setSpacing(12)

        # 左侧：同步文件夹文件
        left_layout = QVBoxLayout()
        left_label = QLabel("同步文件夹")
        left_label.setStyleSheet("color: #666; font-weight: bold;")
        left_layout.addWidget(left_label)

        self.folder_list = QListWidget()
        self.folder_list.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.folder_list.itemDoubleClicked.connect(self._on_folder_double_click)
        left_layout.addWidget(self.folder_list)

        left_frame = QFrame()
        left_frame.setLayout(left_layout)
        content_layout.addWidget(left_frame, 1)

        # 右侧：最近文件
        right_layout = QVBoxLayout()
        right_label = QLabel("最近打开")
        right_label.setStyleSheet("color: #666; font-weight: bold;")
        right_layout.addWidget(right_label)

        self.recent_list = QListWidget()
        self.recent_list.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.recent_list.itemDoubleClicked.connect(self._on_recent_double_click)
        right_layout.addWidget(self.recent_list)

        right_frame = QFrame()
        right_frame.setLayout(right_layout)
        content_layout.addWidget(right_frame, 1)

        layout.addLayout(content_layout, 1)

        # 底部：不再显示 + 按钮
        bottom_layout = QHBoxLayout()
        self.dont_show_check = QCheckBox("启动时不再显示此窗口")
        bottom_layout.addWidget(self.dont_show_check)
        bottom_layout.addStretch()

        btn_new = QPushButton("新建文档")
        btn_new.clicked.connect(self._on_new_clicked)
        bottom_layout.addWidget(btn_new)

        btn_open = QPushButton("打开其他文件...")
        btn_open.clicked.connect(self._on_open_other_clicked)
        bottom_layout.addWidget(btn_open)

        btn_cancel = QPushButton("取消")
        btn_cancel.clicked.connect(self.reject)
        bottom_layout.addWidget(btn_cancel)

        layout.addLayout(bottom_layout)

        # 加载数据
        self._load_folder_files(sync_folder)
        self._load_recent_files()

        # 默认选中第一个存在的列表
        if self.folder_list.count() > 0:
            self.folder_list.setCurrentRow(0)
        elif self.recent_list.count() > 0:
            self.recent_list.setCurrentRow(0)

    def _load_folder_files(self, sync_folder: str):
        if not sync_folder or not os.path.exists(sync_folder):
            item = QListWidgetItem("（未配置同步文件夹）")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.folder_list.addItem(item)
            return

        try:
            files = []
            for name in os.listdir(sync_folder):
                full_path = os.path.join(sync_folder, name)
                if os.path.isfile(full_path) and name.lower().endswith(('.md', '.txt', '.text', '.json', '.yaml', '.yml')):
                    mtime = os.path.getmtime(full_path)
                    files.append((name, full_path, mtime))
            files.sort(key=lambda x: x[2], reverse=True)

            if not files:
                item = QListWidgetItem("（文件夹为空）")
                item.setFlags(Qt.ItemFlag.NoItemFlags)
                self.folder_list.addItem(item)
                return

            for name, full_path, mtime in files[:30]:
                item = QListWidgetItem(name)
                item.setData(Qt.ItemDataRole.UserRole, full_path)
                item.setToolTip(full_path)
                self.folder_list.addItem(item)
        except Exception as e:
            item = QListWidgetItem(f"（读取失败：{str(e)}）")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.folder_list.addItem(item)

    def _load_recent_files(self):
        recent = get_recent_files()
        if not recent:
            item = QListWidgetItem("（暂无记录）")
            item.setFlags(Qt.ItemFlag.NoItemFlags)
            self.recent_list.addItem(item)
            return

        for r in recent:
            path = r.get("path", "")
            name = r.get("name", os.path.basename(path))
            exists = os.path.exists(path)
            item = QListWidgetItem(name)
            item.setData(Qt.ItemDataRole.UserRole, path)
            item.setToolTip(path)
            if not exists:
                item.setForeground(QColor("#999"))
                item.setText(f"{name} （已不存在）")
            self.recent_list.addItem(item)

    def _on_folder_double_click(self, item):
        path = item.data(Qt.ItemDataRole.UserRole)
        if path and os.path.isfile(path):
            self.selected_path = path
            self.accept()

    def _on_recent_double_click(self, item):
        path = item.data(Qt.ItemDataRole.UserRole)
        if path and os.path.isfile(path):
            self.selected_path = path
            self.accept()

    def _on_new_clicked(self):
        self.selected_path = None
        self.done(2)

    def _on_open_other_clicked(self):
        self.selected_path = None
        self.done(3)

    def dont_show_again(self) -> bool:
        return self.dont_show_check.isChecked()


# ==================== Markdown 语法高亮器（优化版） ====================
class MarkdownHighlighter(QSyntaxHighlighter):
    # 类常量：限制高亮的最大行数，避免大文件卡顿
    MAX_HIGHLIGHT_LINES = 1000

    def __init__(self, parent=None):
        super().__init__(parent)
        self.rules = []
        # 预编译正则表达式以提升性能
        self._compile_patterns()

    def _compile_patterns(self):
        # 标题 (# ...)
        title_format = QTextCharFormat()
        title_format.setForeground(QColor(0, 120, 200))
        title_format.setFontWeight(QFont.Weight.Bold)
        self.rules.append((re.compile(r"^#{1,6}\s+.*"), title_format))

        # 粗体 (**bold**)
        bold_format = QTextCharFormat()
        bold_format.setFontWeight(QFont.Weight.Bold)
        self.rules.append((re.compile(r"\*\*[^*]+\*\*"), bold_format))

        # 斜体 (*italic*)
        italic_format = QTextCharFormat()
        italic_format.setFontItalic(True)
        self.rules.append((re.compile(r"(?<!\*)\*[^*]+\*(?!\*)"), italic_format))

        # 行内代码 (`code`)
        code_format = QTextCharFormat()
        code_format.setForeground(QColor(150, 100, 50))
        code_format.setFont(QFont("Consolas"))
        self.rules.append((re.compile(r"`[^`]+`"), code_format))

        # 链接 [text](url)
        link_format = QTextCharFormat()
        link_format.setForeground(QColor(0, 150, 0))
        link_format.setFontUnderline(True)
        self.rules.append((re.compile(r"\[[^\]]+\]\([^\)]+\)"), link_format))

    def highlightBlock(self, text):
        # 限制高亮处理的行数，避免大文件卡顿
        block_number = self.currentBlock().blockNumber()
        if block_number >= self.MAX_HIGHLIGHT_LINES:
            return

        # 先清除格式再应用
        self.setFormat(0, len(text), QTextCharFormat())
        for pattern, fmt in self.rules:
            for match in pattern.finditer(text):
                start, end = match.span()
                self.setFormat(start, end - start, fmt)


# ==================== 自定义编辑器 ====================
class CodeEditor(QTextEdit):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.file_path = None
        self.highlighter = None
        self._is_loading = False  # 加载标志，避免触发预览更新
        self.setFont(QFont("Consolas", 10))
        self.setTabStopDistance(4 * self.fontMetrics().horizontalAdvance(' '))

    def set_markdown_highlighting(self, enabled):
        if enabled:
            if self.highlighter is None:
                self.highlighter = MarkdownHighlighter(self.document())
            return
        if self.highlighter is not None:
            self.highlighter.setDocument(None)
            self.highlighter = None

    @staticmethod
    def is_markdown_path(path):
        ext = QFileInfo(path).suffix().lower()
        return ext in ["md", "markdown"]

    def set_file_path(self, path):
        self.file_path = path

    def load_file(self, path):
        try:
            content = self._read_file_with_fallback(path)
            self.set_markdown_highlighting(False)
            self._is_loading = True
            cursor = self.textCursor()
            cursor.beginEditBlock()
            cursor.select(QTextCursor.SelectionType.Document)
            cursor.removeSelectedText()
            cursor.insertText(content)
            cursor.endEditBlock()
            self.file_path = path
            self._is_loading = False
            QTimer.singleShot(50, lambda: self.set_markdown_highlighting(self.is_markdown_path(path)))
            self.document().setModified(False)
            return True
        except Exception as e:
            self._is_loading = False
            QMessageBox.critical(self, "错误", f"无法打开文件：{str(e)}")
            return False

    @staticmethod
    def _read_file_with_fallback(path):
        with open(path, 'rb') as f:
            raw = f.read()
        if raw.startswith(b'\xff\xfe'):
            try:
                return raw.decode('utf-16-le')
            except Exception:
                pass
        elif raw.startswith(b'\xfe\xff'):
            try:
                return raw.decode('utf-16-be')
            except Exception:
                pass
        elif raw.startswith(b'\xef\xbb\xbf'):
            try:
                return raw.decode('utf-8-sig')
            except Exception:
                pass
        if len(raw) >= 2:
            null_count = raw.count(b'\x00')
            if null_count > len(raw) * 0.2:
                try:
                    return raw.decode('utf-16-le')
                except Exception:
                    try:
                        return raw.decode('utf-16-be')
                    except Exception:
                        pass
        encodings = ['utf-8', 'utf-8-sig', 'utf-16', 'utf-16-le', 'gbk', 'gb2312', 'gb18030', 'latin-1']
        for encoding in encodings:
            try:
                return raw.decode(encoding)
            except (UnicodeDecodeError, UnicodeError):
                continue
        return raw.decode('utf-8', errors='replace')

    def save_to_path(self, path):
        try:
            with open(path, 'w', encoding='utf-8') as f:
                f.write(self.toPlainText())
            self.file_path = path
            self.set_markdown_highlighting(self.is_markdown_path(path))
            self.document().setModified(False)
            return True
        except Exception as e:
            QMessageBox.critical(self, "错误", f"无法保存文件：{str(e)}")
            return False


# ==================== 悬停放大容器（产品亮点，保留） ====================
class HoverWidget(QWidget):
    def __init__(self, widget, on_enter, on_leave, parent=None):
        super().__init__(parent)
        self.on_enter_callback = on_enter
        self.on_leave_callback = on_leave
        layout = QVBoxLayout()
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(widget)
        self.setLayout(layout)
        self.widget = widget
        self.setAttribute(Qt.WidgetAttribute.WA_StyledBackground, True)

    def enterEvent(self, event):
        self.on_enter_callback(self)
        super().enterEvent(event)

    def leaveEvent(self, event):
        self.on_leave_callback(self)
        super().leaveEvent(event)


class PreviewBridge(QObject):
    scrollPercentChanged = pyqtSignal(float)

    @pyqtSlot(float)
    def reportScrollPercent(self, percent):
        self.scrollPercentChanged.emit(percent)


def get_resource_base_dir():
    return getattr(sys, "_MEIPASS", os.path.dirname(__file__))


def get_resource_url(*parts):
    return QUrl.fromLocalFile(os.path.join(get_resource_base_dir(), *parts))


def convert_special_markdown(text):
    """
    简化版 Markdown 扩展处理
    移除了 MathJax 和 Mermaid 的特殊处理
    仅保留下标/上标转换
    """
    # 下标 ~text~
    text = re.sub(r"(?<!\\)~([^~\n]+?)~", r"<sub>\1</sub>", text)
    # 上标 ^text^
    text = re.sub(r"(?<!\\)\^([^\^\n]+?)\^", r"<sup>\1</sup>", text)
    return text


def build_preview_html(markdown_text):
    """
    优化版 HTML 生成
    移除 MathJax 和 Mermaid，仅保留基础 Markdown 渲染
    """
    markdown_module = get_markdown_module()
    converted_text = convert_special_markdown(markdown_text)

    # 简化扩展：移除 codehilite（太重），保留基础功能
    body_html = markdown_module.markdown(
        converted_text,
        extensions=['extra', 'tables', 'fenced_code']
    )

    # 精简的 CSS 样式
    css = """
    <style>
        :root { color-scheme: light; }
        html, body {
            margin: 0; padding: 0;
            background: #f7f7fb;
            color: #1f2937;
            font-family: 'Segoe UI', 'Noto Sans SC', sans-serif;
            font-size: 14px; line-height: 1.75;
        }
        #content {
            max-width: 980px; margin: 0 auto;
            padding: 24px 28px 72px; box-sizing: border-box;
        }
        h1, h2, h3, h4, h5, h6 {
            line-height: 1.25; margin: 1.2em 0 0.6em; color: #0f172a;
        }
        h1 { font-size: 2rem; border-bottom: 1px solid #dbe1ea; padding-bottom: 0.3em; }
        h2 { font-size: 1.5rem; border-bottom: 1px solid #e5e7eb; padding-bottom: 0.2em; }
        h3 { font-size: 1.25rem; }
        p, ul, ol, blockquote, table, pre { margin: 0.9em 0; }
        a { color: #0366d6; text-decoration: none; }
        a:hover { text-decoration: underline; }
        code {
            font-family: 'Cascadia Mono', 'Consolas', monospace;
            background: #eef2f7; color: #0f172a;
            border-radius: 6px; padding: 0.15em 0.35em;
        }
        pre {
            background: #0b1020; color: #e5eefb;
            border-radius: 12px; padding: 16px; overflow: auto;
        }
        pre code { background: transparent; color: inherit; padding: 0; }
        blockquote {
            border-left: 4px solid #7c3aed; margin-left: 0;
            padding: 0.4em 1em; color: #4b5563;
            background: rgba(124, 58, 237, 0.06);
            border-radius: 0 10px 10px 0;
        }
        table { border-collapse: collapse; width: 100%; display: block; overflow-x: auto; }
        th, td { border: 1px solid #d1d5db; padding: 0.55em 0.8em; text-align: left; }
        th { background: #eef2ff; }
        img { max-width: 100%; height: auto; }
        hr { border: none; border-top: 1px solid #d1d5db; margin: 1.5em 0; }
    </style>
    """

    # 精简的 JavaScript：仅保留滚动同步，移除 MathJax/Mermaid
    script = """
    <script src="qrc:///qtwebchannel/qwebchannel.js"></script>
    <script>
        window.setScrollPercent = function(percent) {
            const doc = document.scrollingElement || document.documentElement;
            const maxScroll = Math.max(0, doc.scrollHeight - doc.clientHeight);
            const clamped = Math.max(0, Math.min(1, percent));
            doc.scrollTop = maxScroll <= 0 ? 0 : Math.round(maxScroll * clamped);
        };

        document.addEventListener('DOMContentLoaded', function() {
            new QWebChannel(qt.webChannelTransport, function(channel) {
                window.tmlBridge = channel.objects.tmlBridge;
                const sendScroll = function() {
                    const doc = document.scrollingElement || document.documentElement;
                    const maxScroll = Math.max(1, doc.scrollHeight - doc.clientHeight);
                    window.tmlBridge.reportScrollPercent(doc.scrollTop / maxScroll);
                };

                let scheduled = false;
                window.addEventListener('scroll', function() {
                    if (scheduled) return;
                    scheduled = true;
                    requestAnimationFrame(function() {
                        scheduled = false;
                        sendScroll();
                    });
                }, { passive: true });

                // 初始化滚动位置
                if (typeof window.__pendingScrollPercent === 'number') {
                    window.setScrollPercent(window.__pendingScrollPercent);
                }
            });
        });
    </script>
    """

    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
{css}
{script}
</head>
<body>
<div id="content">
{body_html}
</div>
</body>
</html>"""


# ==================== 主窗口 ====================
class MarkdownEditor(QMainWindow):
    def __init__(self, initial_paths=None):
        super().__init__()
        self.startup_paths = list(initial_paths or [])
        log_startup(f"window init begin, paths={len(self.startup_paths)}")
        icon_path = get_icon_path()
        if icon_path:
            self.setWindowIcon(QIcon(icon_path))
        self.setWindowTitle("TML Markdown 编辑器")
        self.resize(700, 600)
        self.setAcceptDrops(True)

        self.split_enabled = False
        self.hovered_side = None
        self._syncing_scroll = False
        self._preview_loaded = False
        self._preview_scroll_percent = 0.0

        # 防抖定时器：避免每次按键都触发预览更新
        self._preview_timer = QTimer()
        self._preview_timer.setSingleShot(True)
        self._preview_timer.timeout.connect(self._do_update_preview)

        # 中央分割器
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        self.setCentralWidget(self.splitter)

        # 左侧预览区（懒初始化）
        self.preview = None
        self.preview_bridge = None
        self.preview_channel = None
        preview_placeholder = QWidget()
        self.left_container = HoverWidget(
            preview_placeholder,
            self.on_hover_enter,
            self.on_hover_leave,
            self.splitter
        )

        # 右侧多标签页编辑区
        self.tab_widget = QTabWidget()
        self.tab_widget.setTabsClosable(True)
        self.tab_widget.tabCloseRequested.connect(self.close_tab)
        self.tab_widget.currentChanged.connect(self.on_current_tab_changed)
        self.right_container = HoverWidget(
            self.tab_widget,
            self.on_hover_enter,
            self.on_hover_leave,
            self.splitter
        )

        self.splitter.addWidget(self.left_container)
        self.splitter.addWidget(self.right_container)

        # 初始单屏模式
        self.left_container.hide()
        self.splitter.setSizes([0, self.width()])

        self.create_menu_bar()

        self.status_bar = QStatusBar()
        self.setStatusBar(self.status_bar)
        self.status_bar.showMessage("就绪")

        # 状态栏同步状态指示
        self.sync_status_label = QLabel("")
        self.sync_status_label.setStyleSheet("color: #888; padding: 0 8px;")
        self.status_bar.addPermanentWidget(self.sync_status_label)

        self.init_sync_client()
        QTimer.singleShot(500, self.refresh_sync_status)

        # 自动同步定时器（每60秒）
        self.auto_sync_timer = QTimer(self)
        self.auto_sync_timer.timeout.connect(self._auto_sync_tick)
        self._is_sync_running = False
        self._current_sync_worker = None
        self._setup_auto_sync()

        # 同步状态旋转动画
        self._sync_spin_angle = 0
        self._sync_spin_timer = QTimer(self)
        self._sync_spin_timer.timeout.connect(self._update_sync_spin)
        self._sync_spin_timer.setInterval(50)

        if self.startup_paths:
            QTimer.singleShot(0, self.open_startup_files)
        else:
            settings = load_settings()
            if settings.get("show_quick_open_on_start", True):
                QTimer.singleShot(100, self._show_quick_open_dialog)
            else:
                self.new_tab()
        log_startup("window init end")

    def _show_quick_open_dialog(self):
        sync_folder = ""
        if hasattr(self, "sync_folder") and self.sync_folder:
            sync_folder = self.sync_folder
        dialog = QuickOpenDialog(sync_folder, self)
        result = dialog.exec()

        if dialog.dont_show_again():
            settings = load_settings()
            settings["show_quick_open_on_start"] = False
            save_settings(settings)

        if result == 1 and dialog.selected_path:
            self.open_file(dialog.selected_path)
        elif result == 2:
            self.new_tab()
        elif result == 3:
            self.open_file()
        else:
            self.new_tab()

    def open_startup_files(self):
        log_startup(f"opening startup files: {len(self.startup_paths)}")
        if not self.open_file_paths(self.startup_paths):
            self.new_tab()
        self.startup_paths = []

    def ensure_preview_initialized(self):
        if self.preview is not None:
            return

        log_startup("initializing preview webengine")
        from PyQt6.QtWebEngineWidgets import QWebEngineView
        from PyQt6.QtWebEngineCore import QWebEngineSettings
        from PyQt6.QtWebChannel import QWebChannel

        self.preview = QWebEngineView()
        self.preview.setContextMenuPolicy(Qt.ContextMenuPolicy.NoContextMenu)
        self.preview.settings().setAttribute(QWebEngineSettings.WebAttribute.LocalContentCanAccessFileUrls, True)
        self.preview.settings().setAttribute(QWebEngineSettings.WebAttribute.LocalContentCanAccessRemoteUrls, False)
        self.preview.page().loadFinished.connect(self.on_preview_load_finished)

        self.preview_bridge = PreviewBridge(self)
        self.preview_bridge.scrollPercentChanged.connect(self.on_preview_scroll_percent)
        self.preview_channel = QWebChannel(self.preview.page())
        self.preview_channel.registerObject("tmlBridge", self.preview_bridge)
        self.preview.page().setWebChannel(self.preview_channel)

        old_left = self.left_container
        self.left_container = HoverWidget(
            self.preview,
            self.on_hover_enter,
            self.on_hover_leave,
            self.splitter
        )
        self.splitter.insertWidget(0, self.left_container)
        old_left.hide()
        old_left.setParent(None)
        old_left.deleteLater()

    # ========== 滚动同步 ==========
    def sync_scroll(self, source_scrollbar, target_scrollbar):
        if self._syncing_scroll or not self.split_enabled:
            return
        self._syncing_scroll = True
        src_min = source_scrollbar.minimum()
        src_max = source_scrollbar.maximum()
        src_val = source_scrollbar.value()
        if src_max > src_min:
            percent = (src_val - src_min) / (src_max - src_min)
            tgt_min = target_scrollbar.minimum()
            tgt_max = target_scrollbar.maximum()
            tgt_val = int(tgt_min + percent * (tgt_max - tgt_min))
            target_scrollbar.setValue(tgt_val)
        self._syncing_scroll = False

    def scrollbar_percent(self, scrollbar):
        minimum = scrollbar.minimum()
        maximum = scrollbar.maximum()
        if maximum <= minimum:
            return 0.0
        return (scrollbar.value() - minimum) / (maximum - minimum)

    def set_scrollbar_percent(self, scrollbar, percent):
        minimum = scrollbar.minimum()
        maximum = scrollbar.maximum()
        if maximum <= minimum:
            scrollbar.setValue(minimum)
            return
        clamped = max(0.0, min(1.0, percent))
        scrollbar.setValue(int(minimum + clamped * (maximum - minimum)))

    def set_preview_scroll_percent(self, percent):
        self._preview_scroll_percent = max(0.0, min(1.0, percent))
        if not self._preview_loaded or not self.split_enabled or self.preview is None:
            return
        self.preview.page().runJavaScript(f"window.setScrollPercent({self._preview_scroll_percent:.8f});")

    def connect_scroll_sync(self):
        editor = self.current_editor()
        if editor and self.split_enabled:
            editor_vscroll = editor.verticalScrollBar()
            try:
                editor_vscroll.valueChanged.disconnect(self.on_editor_scroll)
            except TypeError:
                pass
            editor_vscroll.valueChanged.connect(self.on_editor_scroll)
            self.on_editor_scroll(editor_vscroll.value())

    def disconnect_scroll_sync(self):
        editor = self.current_editor()
        if editor:
            try:
                editor.verticalScrollBar().valueChanged.disconnect(self.on_editor_scroll)
            except TypeError:
                pass

    def on_editor_scroll(self, value):
        if not self.split_enabled or self._syncing_scroll:
            return
        editor = self.current_editor()
        if editor:
            percent = self.scrollbar_percent(editor.verticalScrollBar())
            self.set_preview_scroll_percent(percent)

    def on_preview_scroll_percent(self, percent):
        if not self.split_enabled or self._syncing_scroll:
            return
        editor = self.current_editor()
        if editor:
            current_percent = self.scrollbar_percent(editor.verticalScrollBar())
            if abs(current_percent - percent) < 0.01:
                return
            self._syncing_scroll = True
            self.set_scrollbar_percent(editor.verticalScrollBar(), percent)
            self._syncing_scroll = False

    # ========== 悬停放大（产品亮点，保留） ==========
    def on_hover_enter(self, hover_widget):
        if not self.split_enabled:
            return
        if hover_widget == self.left_container:
            self.hovered_side = "left"
        elif hover_widget == self.right_container:
            self.hovered_side = "right"
        self.adjust_splitter_sizes()

    def on_hover_leave(self, hover_widget):
        if not self.split_enabled:
            return
        self.hovered_side = None
        self.adjust_splitter_sizes()

    def adjust_splitter_sizes(self):
        if not self.split_enabled:
            return
        total = self.splitter.width()
        if total <= 0:
            return
        if self.hovered_side == "left":
            self.splitter.setSizes([int(total * 0.7), int(total * 0.3)])
        elif self.hovered_side == "right":
            self.splitter.setSizes([int(total * 0.3), int(total * 0.7)])
        else:
            self.splitter.setSizes([total // 2, total // 2])

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if self.split_enabled:
            self.adjust_splitter_sizes()

    # ========== 分屏模式 ==========
    def set_split_mode(self, enabled):
        self.split_enabled = enabled
        if enabled:
            self.ensure_preview_initialized()
            self.left_container.show()
            total = self.splitter.width()
            if total > 0:
                self.splitter.setSizes([total // 2, total // 2])
            self.update_preview()
            self.connect_scroll_sync()
            self.status_bar.showMessage("分屏模式已开启")
        else:
            self.disconnect_scroll_sync()
            self.left_container.hide()
            self.splitter.setSizes([0, self.splitter.width()])
            self.status_bar.showMessage("分屏模式已关闭")
        self.split_action.setChecked(enabled)

    def toggle_split_mode(self):
        self.set_split_mode(not self.split_enabled)

    # ========== 多标签页管理 ==========
    def new_tab(self, file_path=None, content=""):
        editor = CodeEditor()
        if file_path:
            if not editor.load_file(file_path):
                return None
            tab_name = os.path.basename(file_path)
        else:
            editor.setPlainText(content)
            tab_name = "未命名"
        index = self.tab_widget.addTab(editor, tab_name)
        self.tab_widget.setCurrentIndex(index)
        editor.textChanged.connect(lambda: self.on_editor_text_changed(editor))
        editor.document().modificationChanged.connect(
            lambda modified: self.update_tab_title(editor, modified)
        )
        if self.split_enabled:
            self.connect_scroll_sync()
        return editor

    def open_file_paths(self, paths):
        valid_paths = [path for path in paths if path and os.path.isfile(path)]
        if not valid_paths:
            return False

        log_startup(f"open_file_paths valid={len(valid_paths)}")
        opened_any = False
        for path in valid_paths:
            log_startup(f"opening file: {os.path.basename(path)}")
            editor = self.new_tab(file_path=path)
            if editor is None:
                log_startup(f"failed file: {path}")
                continue
            opened_any = True
            self.status_bar.showMessage(f"已打开：{path}")
            try:
                add_recent_file(path)
            except Exception:
                pass

        if opened_any:
            current = self.current_editor()
            if current and current.file_path:
                ext = QFileInfo(current.file_path).suffix().lower()
                is_md = ext in ["md", "markdown"]
                if is_md != self.split_enabled:
                    self.set_split_mode(is_md)
                elif self.split_enabled:
                    self.update_preview()

        return opened_any

    def close_tab(self, index):
        editor = self.tab_widget.widget(index)
        if editor and editor.document().isModified():
            tab_name = self.tab_widget.tabText(index)
            ret = QMessageBox.question(
                self, "未保存",
                f"文档「{tab_name}」已修改，是否保存？",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No | QMessageBox.StandardButton.Cancel
            )
            if ret == QMessageBox.StandardButton.Yes:
                if not self.save_current_editor(editor):
                    return
            elif ret == QMessageBox.StandardButton.Cancel:
                return
        self.tab_widget.removeTab(index)
        if self.tab_widget.count() == 0:
            self.new_tab()
        elif self.split_enabled:
            self.connect_scroll_sync()

    def update_tab_title(self, editor, modified):
        index = self.tab_widget.indexOf(editor)
        if index == -1:
            return
        base_name = os.path.basename(editor.file_path) if editor.file_path else "未命名"
        title = base_name + " *" if modified else base_name
        self.tab_widget.setTabText(index, title)

    def on_editor_text_changed(self, editor):
        """使用防抖机制延迟预览更新，避免每次按键都触发"""
        # 如果正在加载文件，不触发预览更新
        if getattr(editor, '_is_loading', False):
            return
        if self.split_enabled and editor == self.current_editor():
            self._preview_timer.start(300)  # 300ms 防抖

    def _do_update_preview(self):
        """实际执行预览更新"""
        editor = self.current_editor()
        if not editor or not self.split_enabled:
            return
        self.ensure_preview_initialized()
        log_startup("rendering preview")
        md_text = editor.toPlainText()
        editor_vscroll = editor.verticalScrollBar()
        self._preview_scroll_percent = self.scrollbar_percent(editor_vscroll)
        self._preview_loaded = False
        try:
            full_html = build_preview_html(md_text)
        except Exception as e:
            err_msg = html_lib.escape(str(e))
            full_html = f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<style>
body {{ background: #f7f7fb; color: #1f2937; font-family: sans-serif; padding: 24px; }}
h1 {{ color: #dc2626; }}
pre {{ background: #fee2e2; padding: 12px; border-radius: 6px; overflow-x: auto; }}
</style>
</head>
<body>
<h1>渲染错误</h1>
<p>Markdown 渲染失败，请检查依赖是否正确安装。</p>
<pre>{err_msg}</pre>
</body>
</html>"""
        self.preview.setHtml(full_html)

    def on_current_tab_changed(self, index):
        editor = self.current_editor()
        if editor:
            path = editor.file_path
            if path:
                self.setWindowTitle(f"TML Markdown 编辑器 - {os.path.basename(path)}")
                ext = QFileInfo(path).suffix().lower()
                is_md = ext in ["md", "markdown"]
                if is_md != self.split_enabled:
                    self.set_split_mode(is_md)
            else:
                self.setWindowTitle("TML Markdown 编辑器")
            if self.split_enabled:
                self.connect_scroll_sync()
            self.status_bar.showMessage(f"当前文件：{path if path else '未保存'}")

    def current_editor(self):
        return self.tab_widget.currentWidget()

    def save_current_editor(self, editor=None):
        if editor is None:
            editor = self.current_editor()
        if not editor:
            return False
        if editor.file_path:
            return editor.save_to_path(editor.file_path)
        else:
            return self.save_as_current_editor(editor)

    def save_as_current_editor(self, editor=None):
        if editor is None:
            editor = self.current_editor()
        if not editor:
            return False
        path, _ = QFileDialog.getSaveFileName(
            self, "保存文件", "",
            "Markdown 文件 (*.md);;文本文件 (*.txt);;LaTeX 文件 (*.tex);;所有文件 (*)"
        )
        if path:
            if editor.save_to_path(path):
                self.update_tab_title(editor, False)
                ext = QFileInfo(path).suffix().lower()
                if ext in ["md", "markdown"] and not self.split_enabled:
                    reply = QMessageBox.question(self, "建议", "是否开启分屏模式以预览 Markdown？",
                                                 QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
                    if reply == QMessageBox.StandardButton.Yes:
                        self.set_split_mode(True)
                return True
        return False

    # ========== 文件操作 ==========
    def new_file(self):
        self.new_tab()

    def open_file(self, path: str = ""):
        if path:
            self.open_file_paths([path])
            return
        path, _ = QFileDialog.getOpenFileName(
            self, "打开文件", "",
            "文本文件 (*.txt *.md *.markdown *.tex);;所有文件 (*)"
        )
        if path:
            self.open_file_paths([path])

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls() and any(url.isLocalFile() for url in event.mimeData().urls()):
            event.acceptProposedAction()
        else:
            super().dragEnterEvent(event)

    def dropEvent(self, event):
        paths = [url.toLocalFile() for url in event.mimeData().urls() if url.isLocalFile()]
        if self.open_file_paths(paths):
            event.acceptProposedAction()
        else:
            super().dropEvent(event)

    def save_file(self):
        editor = self.current_editor()
        if editor:
            if editor.file_path:
                editor.save_to_path(editor.file_path)
                self.update_tab_title(editor, False)
            else:
                self.save_as_current_editor(editor)

    def save_as_file(self):
        self.save_as_current_editor()

    # ========== 预览更新（直接调用，无需防抖） ==========
    def update_preview(self):
        """立即更新预览（用于切换标签等场景）"""
        self._preview_timer.stop()
        self._do_update_preview()

    def on_preview_load_finished(self, ok):
        self._preview_loaded = bool(ok)
        if not ok or not self.split_enabled:
            return
        self.preview.page().runJavaScript(
            f"window.__pendingScrollPercent = {self._preview_scroll_percent:.8f};"
        )

    # ========== 导出 Word ==========
    def export_to_word(self):
        editor = self.current_editor()
        if not editor:
            return
        Document = get_docx_document_class()
        if Document is None:
            QMessageBox.warning(self, "缺少依赖", "未安装 python-docx，无法导出 Word。")
            return
        path, _ = QFileDialog.getSaveFileName(
            self, "导出为 Word", "", "Word 文档 (*.docx)"
        )
        if not path:
            return
        try:
            doc = Document()
            self.add_markdown_to_docx(doc, editor.toPlainText())
            doc.save(path)
            self.status_bar.showMessage(f"已导出：{path}")
        except Exception as e:
            QMessageBox.critical(self, "错误", f"导出失败：{str(e)}")

    def add_markdown_to_docx(self, doc, text):
        import re
        lines = text.splitlines()
        in_code_block = False
        for line in lines:
            stripped = line.rstrip()
            if stripped.startswith("```"):
                in_code_block = not in_code_block
                continue
            if in_code_block:
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(stripped)
                run.font.name = "Consolas"
                continue
            if not stripped:
                doc.add_paragraph("")
                continue
            if re.match(r"^#{1,6}\s+", stripped):
                level = len(stripped.split(" ", 1)[0])
                title = stripped[level + 1:]
                doc.add_heading(title, level=level)
                continue
            if re.match(r"^\d+\.\s+", stripped):
                content = re.sub(r"^\d+\.\s+", "", stripped)
                paragraph = doc.add_paragraph(style="List Number")
                self.add_inline_runs(paragraph, content)
                continue
            if stripped.startswith("- ") or stripped.startswith("* "):
                content = stripped[2:]
                paragraph = doc.add_paragraph(style="List Bullet")
                self.add_inline_runs(paragraph, content)
                continue
            if stripped.startswith("> "):
                paragraph = doc.add_paragraph(style="Intense Quote")
                self.add_inline_runs(paragraph, stripped[2:])
                continue
            paragraph = doc.add_paragraph()
            self.add_inline_runs(paragraph, stripped)

    def add_inline_runs(self, paragraph, text):
        import re
        token_re = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
        pos = 0
        for match in token_re.finditer(text):
            if match.start() > pos:
                paragraph.add_run(text[pos:match.start()])
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
            pos = match.end()
        if pos < len(text):
            paragraph.add_run(text[pos:])

    # ========== 帮助窗口 ==========
    def show_help(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("使用指导")
        dialog.resize(520, 420)

        layout = QVBoxLayout()
        guide = QTextEdit()
        guide.setReadOnly(True)
        guide.setPlainText(
            "使用指导\n"
            "\n"
            "1. 新建/打开\n"
            "- 文件菜单中可新建或打开文件。\n"
            "- 支持.txt, .md, .tex, .py文件。\n"
            "\n"
            "2. 分屏预览\n"
            "- 打开.md类型文件自动进入分屏模式。\n"
            "- 视图菜单可开启或关闭分屏模式。\n"
            "- 分屏时左侧预览，右侧编辑，滚动自动同步。\n"
            "- 鼠标悬停在某侧可自动放大该侧比例。\n"
            "\n"
            "3. 预览缩放\n"
            "- 视图菜单可放大或缩小预览字体。\n"
            "\n"
            "4. 导出 Word\n"
            "- 文件菜单可导出为 .docx。\n"
            "\n"
            "本项目旨在提供一个方便、轻量的小窗口Markdown编辑体验\n"
            "欢迎反馈和建议！\n"
            "问题提交：https://github.com/Loryage\n"
        )
        layout.addWidget(guide)

        close_btn = QPushButton("关闭")
        close_btn.clicked.connect(dialog.accept)
        layout.addWidget(close_btn)

        dialog.setLayout(layout)
        dialog.exec()

    # ========== 同步功能 ==========
    def init_sync_client(self):
        try:
            from sync_client import SyncClient, get_sync_folder
            self.sync_client = SyncClient()
            self.sync_folder = get_sync_folder()
        except Exception as e:
            self.sync_client = None
            self.sync_folder = None

    def sync_now(self):
        if not hasattr(self, "sync_client") or self.sync_client is None:
            self.init_sync_client()
        if not self.sync_client.is_configured():
            QMessageBox.information(self, "同步设置", "请先配置服务器信息")
            self.open_sync_settings()
            return
        self.run_sync()

    def open_sync_settings(self):
        if not hasattr(self, "sync_client") or self.sync_client is None:
            self.init_sync_client()
        dialog = QDialog(self)
        dialog.setWindowTitle("同步设置")
        dialog.setFixedWidth(420)
        layout = QFormLayout()

        server_edit = QLineEdit(self.sync_client.server_url)
        api_key_edit = QLineEdit(self.sync_client.api_key)
        api_key_edit.setEchoMode(QLineEdit.EchoMode.Password)
        folder_label = QLabel(self.sync_folder)

        auto_sync_check = QCheckBox("每60秒自动同步一次")
        auto_sync_check.setChecked(self.sync_client.config.get("auto_sync", False))

        layout.addRow("服务器地址：", server_edit)
        layout.addRow("API Key：", api_key_edit)
        layout.addRow("同步文件夹：", folder_label)
        layout.addRow("自动同步：", auto_sync_check)

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addRow(buttons)

        dialog.setLayout(layout)

        if dialog.exec() == QDialog.DialogCode.Accepted:
            self.sync_client.config["server_url"] = server_edit.text().strip()
            self.sync_client.config["api_key"] = api_key_edit.text().strip()
            self.sync_client.config["auto_sync"] = auto_sync_check.isChecked()
            from sync_client import save_sync_config
            save_sync_config(self.sync_client.config)
            QMessageBox.information(self, "设置已保存", "同步设置已保存")
            self._setup_auto_sync()
            self.refresh_sync_status()

    def open_sync_folder(self):
        if not hasattr(self, "sync_client") or self.sync_client is None:
            self.init_sync_client()
        folder = self.sync_folder
        if folder and os.path.exists(folder):
            try:
                if os.name == "nt":
                    os.startfile(folder)
                else:
                    import subprocess
                    subprocess.Popen(["xdg-open", folder])
            except Exception as e:
                QMessageBox.warning(self, "错误", f"无法打开文件夹：{str(e)}")
        else:
            QMessageBox.information(self, "提示", "同步文件夹不存在")

    def refresh_sync_status(self):
        if not hasattr(self, "sync_client") or self.sync_client is None:
            return
        if self._is_sync_running:
            return
        if self.sync_client.is_configured():
            last_sync = self.sync_client.config.get("last_sync_time", 0)
            auto_on = self.sync_client.config.get("auto_sync", False)
            auto_tag = " | 自动同步已开启" if auto_on else ""
            if last_sync > 0:
                import datetime
                dt = datetime.datetime.fromtimestamp(last_sync / 1000)
                self.status_bar.showMessage(f"同步已配置 | 上次同步：{dt.strftime('%Y-%m-%d %H:%M')}{auto_tag}")
            else:
                self.status_bar.showMessage(f"同步已配置 | 尚未同步{auto_tag}")
            self.sync_status_label.setText("🔄 同步就绪")
            self.sync_status_label.setStyleSheet("color: #888; padding: 0 8px;")
        else:
            self.status_bar.showMessage("同步未配置")
            self.sync_status_label.setText("")

    def _start_sync_indicator(self):
        self._is_sync_running = True
        self._sync_spin_angle = 0
        self._sync_spin_timer.start()
        self.sync_status_label.setStyleSheet("color: #2196F3; padding: 0 8px; font-weight: bold;")
        self._update_sync_spin()

    def _stop_sync_indicator(self, success=True):
        self._is_sync_running = False
        self._sync_spin_timer.stop()
        if success:
            self.sync_status_label.setText("✅ 同步完成")
            self.sync_status_label.setStyleSheet("color: #4CAF50; padding: 0 8px;")
        else:
            self.sync_status_label.setText("❌ 同步失败")
            self.sync_status_label.setStyleSheet("color: #f44336; padding: 0 8px;")
        QTimer.singleShot(3000, self._reset_sync_label)

    def _reset_sync_label(self):
        if not self._is_sync_running:
            self.refresh_sync_status()

    def _update_sync_spin(self):
        self._sync_spin_angle = (self._sync_spin_angle + 15) % 360
        chars = ["⠋", "⠙", "⠹", "⠸", "⠼", "⠴", "⠦", "⠧", "⠇", "⠏"]
        idx = (self._sync_spin_angle // 36) % len(chars)
        self.sync_status_label.setText(f"{chars[idx]} 同步中...")

    def _setup_auto_sync(self):
        """根据配置启动或停止自动同步定时器"""
        if not hasattr(self, "sync_client") or self.sync_client is None:
            return
        auto_on = self.sync_client.config.get("auto_sync", False)
        if auto_on and self.sync_client.is_configured():
            self.auto_sync_timer.start(60000)
        else:
            self.auto_sync_timer.stop()

    def _auto_sync_tick(self):
        """自动同步触发：静默执行，不弹对话框"""
        if self._is_sync_running:
            return
        if not hasattr(self, "sync_client") or self.sync_client is None:
            return
        if not self.sync_client.is_configured():
            return
        self._run_sync_worker(show_dialog=False)

    def _run_sync_worker(self, show_dialog=True):
        """统一的同步执行入口，管理 worker 生命周期"""
        if self._is_sync_running:
            return
        if self._current_sync_worker is not None:
            return

        try:
            self._start_sync_indicator()
            worker = SyncWorker(self.sync_client)
            self._current_sync_worker = worker

            progress_dialog = None
            status_label = None
            progress_bar = None

            if show_dialog:
                progress_dialog = QDialog(self)
                progress_dialog.setWindowTitle("正在同步")
                progress_dialog.setFixedSize(360, 120)
                layout = QVBoxLayout()
                status_label = QLabel("准备同步...")
                progress_bar = QProgressBar()
                progress_bar.setRange(0, 100)
                progress_bar.setValue(0)
                layout.addWidget(status_label)
                layout.addWidget(progress_bar)
                progress_dialog.setLayout(layout)
                progress_dialog.show()

            def update_progress(message, percent):
                try:
                    if status_label and progress_bar and progress_dialog and progress_dialog.isVisible():
                        status_label.setText(message)
                        progress_bar.setValue(percent)
                except Exception:
                    pass

            def on_finished(result):
                try:
                    self._stop_sync_indicator(success=result.get("success", False))
                    if progress_dialog and progress_dialog.isVisible():
                        progress_dialog.accept()
                    if result.get("success"):
                        if show_dialog:
                            QMessageBox.information(self, "同步完成", result.get("message", "同步成功"))
                        else:
                            self.status_bar.showMessage(result.get("message", "自动同步完成"), 5000)
                        self.refresh_sync_status()
                    else:
                        if show_dialog:
                            QMessageBox.warning(self, "同步失败", result.get("message", "未知错误"))
                        else:
                            self.status_bar.showMessage(f"自动同步失败：{result.get('message', '未知错误')}", 5000)
                except Exception as e:
                    log_startup(f"sync on_finished error: {e}")
                finally:
                    QTimer.singleShot(100, self._cleanup_sync_worker)

            worker.progress_signal.connect(update_progress)
            worker.finished_signal.connect(on_finished)
            worker.start()
        except Exception as e:
            log_startup(f"start sync error: {e}")
            self._is_sync_running = False
            self._current_sync_worker = None
            self._sync_spin_timer.stop()
            self.refresh_sync_status()

    def _cleanup_sync_worker(self):
        """安全清理 worker 引用"""
        try:
            if self._current_sync_worker is not None:
                self._current_sync_worker.deleteLater()
                self._current_sync_worker = None
        except Exception as e:
            log_startup(f"cleanup sync worker error: {e}")
            self._current_sync_worker = None

    def run_sync(self):
        self._run_sync_worker(show_dialog=True)

    # ========== 界面构建 ==========
    def create_menu_bar(self):
        menubar = self.menuBar()
        file_menu = menubar.addMenu("文件(&F)")
        new_action = QAction("新建(&N)", self)
        new_action.triggered.connect(self.new_file)
        new_action.setShortcut("Ctrl+N")
        file_menu.addAction(new_action)

        open_action = QAction("打开(&O)...", self)
        open_action.triggered.connect(self.open_file)
        open_action.setShortcut("Ctrl+O")
        file_menu.addAction(open_action)

        save_action = QAction("保存(&S)", self)
        save_action.triggered.connect(self.save_file)
        save_action.setShortcut("Ctrl+S")
        file_menu.addAction(save_action)

        save_as_action = QAction("另存为(&A)...", self)
        save_as_action.triggered.connect(self.save_as_file)
        file_menu.addAction(save_as_action)

        export_action = QAction("导出为 Word(&W)...", self)
        export_action.triggered.connect(self.export_to_word)
        file_menu.addAction(export_action)

        file_menu.addSeparator()
        exit_action = QAction("退出(&X)", self)
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)

        view_menu = menubar.addMenu("视图(&V)")
        zoom_in_action = QAction("放大预览(&+)", self)
        zoom_in_action.triggered.connect(lambda: self.preview_zoom_in())
        view_menu.addAction(zoom_in_action)

        zoom_out_action = QAction("缩小预览(&-)", self)
        zoom_out_action.triggered.connect(lambda: self.preview_zoom_out())
        view_menu.addAction(zoom_out_action)

        view_menu.addSeparator()
        self.split_action = QAction("分屏模式(&P)", self)
        self.split_action.setCheckable(True)
        self.split_action.setChecked(False)
        self.split_action.triggered.connect(self.toggle_split_mode)
        view_menu.addAction(self.split_action)

        sync_menu = menubar.addMenu("同步(&S)")
        sync_now_action = QAction("立即同步(&Y)", self)
        sync_now_action.triggered.connect(self.sync_now)
        sync_now_action.setShortcut("Ctrl+Shift+S")
        sync_menu.addAction(sync_now_action)

        sync_menu.addSeparator()

        sync_folder_action = QAction("打开同步文件夹(&O)", self)
        sync_folder_action.triggered.connect(self.open_sync_folder)
        sync_menu.addAction(sync_folder_action)

        sync_settings_action = QAction("同步设置(&T)...", self)
        sync_settings_action.triggered.connect(self.open_sync_settings)
        sync_menu.addAction(sync_settings_action)

        settings_menu = menubar.addMenu("设置(&E)")
        quick_start_action = QAction("启动时显示快速打开窗口", self)
        quick_start_action.setCheckable(True)
        quick_start_action.setChecked(load_settings().get("show_quick_open_on_start", True))
        quick_start_action.triggered.connect(self._toggle_quick_open_start)
        settings_menu.addAction(quick_start_action)

        help_menu = menubar.addMenu("帮助(&H)")
        help_action = QAction("使用指导(&G)", self)
        help_action.triggered.connect(self.show_help)
        help_menu.addAction(help_action)

    def preview_zoom_in(self):
        if self.split_enabled and self.preview is not None:
            self.preview.setZoomFactor(self.preview.zoomFactor() + 0.1)

    def preview_zoom_out(self):
        if self.split_enabled and self.preview is not None:
            self.preview.setZoomFactor(max(0.4, self.preview.zoomFactor() - 0.1))

    def _toggle_quick_open_start(self, checked):
        settings = load_settings()
        settings["show_quick_open_on_start"] = checked
        save_settings(settings)


def get_icon_path():
    base_dir = getattr(sys, "_MEIPASS", os.path.dirname(__file__))
    icon_path = os.path.join(base_dir, "tml.ico")
    return icon_path if os.path.exists(icon_path) else None


def set_windows_app_id(app_id):
    try:
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(app_id)
    except Exception:
        pass


if __name__ == "__main__":
    log_startup("app entry")
    install_global_exception_hook()
    set_windows_app_id("TML.TMLEditor")
    log_startup("set app id done")

    try:
        import pyi_splash
        _has_pyi_splash = True
    except ImportError:
        _has_pyi_splash = False

    QApplication.setAttribute(Qt.ApplicationAttribute.AA_ShareOpenGLContexts, True)
    log_startup("AA_ShareOpenGLContexts enabled")
    app = QApplication(sys.argv)
    log_startup("QApplication created")
    icon_path = get_icon_path()
    if icon_path:
        app.setWindowIcon(QIcon(icon_path))
    log_startup("window icon applied" if icon_path else "window icon skipped")
    initial_paths = [path for path in sys.argv[1:] if os.path.isfile(path)]
    log_startup(f"argv files={len(initial_paths)}")

    splash = SplashScreen()
    splash.show()
    app.processEvents()
    log_startup("splash shown")
    splash.set_progress(10, "正在初始化界面")
    app.processEvents()

    window = None

    def close_pyi_splash():
        global _has_pyi_splash
        if _has_pyi_splash:
            try:
                import pyi_splash
                if pyi_splash.is_alive():
                    pyi_splash.close()
                    log_startup("pyi_splash closed")
            except Exception:
                pass
            _has_pyi_splash = False

    def finish_startup():
        global window
        splash.set_progress(100, "启动完成")
        app.processEvents()
        close_pyi_splash()
        QTimer.singleShot(200, _final_close)

    def _final_close():
        global window
        try:
            splash.finish(window)
        except Exception:
            splash.close()
        log_startup("splash finished")

    def step_create_window():
        global window
        splash.set_progress(40, "创建主窗口")
        app.processEvents()
        window = MarkdownEditor(initial_paths)
        log_startup("window created")
        splash.set_progress(75, "加载组件")
        app.processEvents()
        QTimer.singleShot(10, step_show_window)

    def step_show_window():
        splash.set_progress(95, "准备就绪")
        app.processEvents()
        window.show()
        window.raise_()
        window.activateWindow()
        log_startup("window shown")
        QTimer.singleShot(50, finish_startup)

    QTimer.singleShot(10, step_create_window)
    sys.exit(app.exec())
